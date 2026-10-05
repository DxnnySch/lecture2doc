"""Write the per-slide transcript Word document."""
from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt

MAX_WORDS = 80


@dataclass
class SpanText:
    slide: int
    start: float
    text: str
    raw: str


def _clean_xml(s: str) -> str:
    """Remove XML-invalid control characters [\x00-\x08\x0b\x0c\x0e-\x1f]."""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', s)


def fmt_time(seconds: float) -> str:
    s = int(seconds)
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    return f"{h}:{m:02d}:{sec:02d}" if h else f"{m:02d}:{sec:02d}"


def paragraphize(text: str, per: int = 4) -> list[str]:
    sentences = [s for s in re.split(r"(?<=[.!?])\s+", text.strip()) if s]
    result = []
    current_group = []

    for s in sentences:
        words = s.split()
        if len(words) > MAX_WORDS:
            # Flush current group first
            if current_group:
                result.append(" ".join(current_group))
                current_group = []
            # Emit long sentence as chunks
            for i in range(0, len(words), MAX_WORDS):
                result.append(" ".join(words[i:i + MAX_WORDS]))
        else:
            # Add to current group
            current_group.append(s)
            if len(current_group) == per:
                result.append(" ".join(current_group))
                current_group = []

    # Flush any remaining group
    if current_group:
        result.append(" ".join(current_group))

    return result


def write_docx(path, title: str, slides, span_texts: list[SpanText],
               thumbnails: bool = True, keep_raw: bool = False) -> None:
    doc = Document()
    doc.add_heading(_clean_xml(title), 0)
    by_slide: dict[int, list[SpanText]] = {}
    for st in sorted(span_texts, key=lambda s: s.start):
        if st.text.strip():
            by_slide.setdefault(st.slide, []).append(st)
    for s in slides:
        heading = f"Slide {s.number} – {_clean_xml(s.title)}" if s.title else f"Slide {s.number}"
        doc.add_heading(heading, 1)
        if thumbnails:
            buf = BytesIO()
            s.image.save(buf, "PNG")
            buf.seek(0)
            doc.add_picture(buf, width=Inches(3.5))
        entries = by_slide.get(s.number, [])
        if not entries:
            doc.add_paragraph().add_run("(no speech)").italic = True
            continue
        timed = len(entries) > 1
        for e in entries:
            for k, para in enumerate(paragraphize(e.text)):
                prefix = f"[{fmt_time(e.start)}] " if timed and k == 0 else ""
                doc.add_paragraph(prefix + _clean_xml(para))
            if keep_raw:
                run = doc.add_paragraph().add_run("Raw: " + _clean_xml(e.raw))
                run.italic = True
                run.font.size = Pt(9)
    doc.save(path)
