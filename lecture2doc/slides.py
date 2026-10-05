"""Slide material: page images from the PDF, titles/text from the PPTX."""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass

import pymupdf
from PIL import Image
from pptx import Presentation

log = logging.getLogger(__name__)

TITLE_LIMIT = 100
TITLE_REGION = 0.6  # titles sit in the upper part of a slide; footers and page numbers sit at the bottom


@dataclass
class Slide:
    number: int
    title: str | None
    text: str
    image: Image.Image


def _squash(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def _first_line(text: str) -> str | None:
    for line in text.splitlines():
        if line.strip():
            return _squash(line)
    return None


def _truncate(line: str | None, limit: int = TITLE_LIMIT) -> str | None:
    if line is not None and len(line) > limit:
        return line[: limit - 1].rstrip() + "…"
    return line


def _fallback_title(body: str, limit: int = TITLE_LIMIT) -> str | None:
    return _truncate(_first_line(body), limit)


def _page_title(page) -> str | None:
    """The largest text in the upper part of the page (topmost on ties); wrapped lines of the same block that
    share that size are joined. Reading order is useless here: in many exports the footer comes first."""
    limit = page.rect.height * TITLE_REGION
    lines = []  # (font size, top y, block number, text)
    for number, block in enumerate(page.get_text("dict")["blocks"]):
        if block.get("type") != 0:
            continue
        for line in block["lines"]:
            text = _squash("".join(span["text"] for span in line["spans"]))
            if text and line["bbox"][1] < limit:
                lines.append((max(span["size"] for span in line["spans"]), line["bbox"][1], number, text))
    if not lines:
        return None
    size, _, block_number, _ = max(lines, key=lambda l: (round(l[0], 1), -l[1]))
    parts = sorted((l for l in lines if l[2] == block_number and abs(l[0] - size) < 0.6), key=lambda l: l[1])
    return _truncate(_squash(" ".join(l[3] for l in parts)))


def pdf_titles(pdf_path) -> list[str | None]:
    with pymupdf.open(pdf_path) as doc:
        return [_page_title(page) for page in doc]


def render_pdf(pdf_path, width: int = 480) -> tuple[list[Image.Image], list[str]]:
    images, texts = [], []
    with pymupdf.open(pdf_path) as doc:
        for page in doc:
            zoom = width / page.rect.width
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            images.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
            texts.append(page.get_text())
    return images, texts


def read_pptx(pptx_path) -> list[tuple[str | None, str]]:
    """(title, body text) for every visible slide; hidden slides are not in the PDF export."""
    out = []
    for s in Presentation(pptx_path).slides:
        if s._element.get("show") == "0":
            continue
        t = s.shapes.title
        title = _squash(t.text_frame.text) if t is not None and t.has_text_frame else ""
        body = "\n".join(
            sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()
        )
        out.append((title or _fallback_title(body), body))
    return out


def load_slides(pdf_path, pptx_path) -> list[Slide]:
    images, pdf_texts = render_pdf(pdf_path)
    try:
        pptx = read_pptx(pptx_path)
    except Exception as e:  # corrupt zip / bad CRC in an embedded media part, etc.
        log.warning("cannot read PPTX (%s: %s); using PDF text for titles", type(e).__name__, e)
        pptx = None
    if pptx is not None and len(pptx) == len(images):
        return [
            Slide(i + 1, title, body or pdf_texts[i].strip(), img)
            for i, ((title, body), img) in enumerate(zip(pptx, images))
        ]
    if pptx is not None:
        log.warning(
            "PPTX has %d visible slides but the PDF has %d pages: slide count does not match; "
            "using PDF text for titles", len(pptx), len(images),
        )
    titles = pdf_titles(pdf_path)
    return [Slide(i + 1, titles[i], t.strip(), img) for i, (t, img) in enumerate(zip(pdf_texts, images))]
