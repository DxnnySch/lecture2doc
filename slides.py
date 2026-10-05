"""Slide material: page images from the PDF, titles/text from the PPTX."""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass

import pymupdf
from PIL import Image
from pptx import Presentation

log = logging.getLogger(__name__)


@dataclass
class Slide:
    number: int
    title: str | None
    text: str
    image: Image.Image


def _squash(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def _first_line(text: str) -> str | None:
    for line in text.splitlines():
        if line.strip():
            return _squash(line)
    return None


def _fallback_title(body: str, limit: int = 100) -> str | None:
    line = _first_line(body)
    if line is not None and len(line) > limit:
        line = line[: limit - 1].rstrip() + "…"
    return line


def render_pdf(pdf_path, width: int = 480) -> tuple[list[Image.Image], list[str]]:
    images, texts = [], []
    with pymupdf.open(pdf_path) as doc:
        for page in doc:
            zoom = width / page.rect.width
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            images.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
            texts.append(page.get_text())
    return images, texts


def read_pptx(pptx_path) -> list[tuple[str | None, str]]:
    """(title, body text) for every visible slide; hidden slides are not in the PDF export."""
    out = []
    for s in Presentation(pptx_path).slides:
        if s._element.get("show") == "0":
            continue
        t = s.shapes.title
        title = _squash(t.text_frame.text) if t is not None and t.has_text_frame else ""
        body = "\n".join(
            sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()
        )
        out.append((title or _fallback_title(body), body))
    return out


def load_slides(pdf_path, pptx_path) -> list[Slide]:
    images, pdf_texts = render_pdf(pdf_path)
    try:
        pptx = read_pptx(pptx_path)
    except Exception as e:  # corrupt zip / bad CRC in an embedded media part, etc.
        log.warning("cannot read PPTX (%s: %s); using PDF text for titles", type(e).__name__, e)
        pptx = None
    if pptx is not None and len(pptx) == len(images):
        return [
            Slide(i + 1, title, body or pdf_texts[i].strip(), img)
            for i, ((title, body), img) in enumerate(zip(pptx, images))
        ]
    if pptx is not None:
        log.warning(
            "PPTX has %d visible slides but the PDF has %d pages: slide count does not match; "
            "using PDF text for titles", len(pptx), len(images),
        )
    return [Slide(i + 1, _first_line(t), t.strip(), img) for i, (t, img) in enumerate(zip(pdf_texts, images))]
