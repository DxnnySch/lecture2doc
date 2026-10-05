from pathlib import Path

import numpy as np
import pymupdf
from pptx import Presentation


def make_pdf(path: Path, n: int, titles: list[str] | None = None) -> Path:
    doc = pymupdf.open()
    for i in range(n):
        page = doc.new_page(width=960, height=540)
        rng = np.random.RandomState(i)
        for _ in range(12):
            x, y = int(rng.randint(0, 800)), int(rng.randint(60, 420))
            w, h = int(rng.randint(40, 160)), int(rng.randint(20, 100))
            page.draw_rect(pymupdf.Rect(x, y, x + w, y + h), color=None, fill=tuple(float(c) for c in rng.rand(3)))
        page.insert_text((40, 40), titles[i] if titles else f"Title {i + 1}", fontsize=24)
    doc.save(path)
    return path


def make_pptx(path: Path, slides: list[tuple[str | None, str, bool]]) -> Path:
    prs = Presentation()
    for title, body, hidden in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        if title is None:
            el = s.shapes.title._element
            el.getparent().remove(el)
        else:
            s.shapes.title.text = title
        s.placeholders[1].text = body
        if hidden:
            s._element.set("show", "0")
    prs.save(path)
    return path
