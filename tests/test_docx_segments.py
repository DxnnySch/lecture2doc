"""The default layout: one page per stretch of the lecture on one slide, in lecture order."""
from docx import Document
from PIL import Image

from docx_out import SpanText, write_docx
from slides import Slide


def deck():
    img = Image.new("RGB", (48, 27), "white")
    return [Slide(1, "Intro", "", img), Slide(2, None, "", img), Slide(3, "Outro", "", img)]


def paragraphs(path):
    return [(p.style.name, p.text, p.paragraph_format.page_break_before) for p in Document(path).paragraphs]


def headings(path):
    return [(t, brk) for style, t, brk in paragraphs(path) if style == "Heading 1"]


def test_every_visit_gets_its_own_page_in_lecture_order(tmp_path):
    out = tmp_path / "o.docx"
    spans = [SpanText(1, 0, "First visit.", "r1", end=10), SpanText(2, 10, "Middle.", "r2", end=20),
             SpanText(1, 20, "Back again.", "r3", end=30), SpanText(3, 30, "The end.", "r4", end=40)]
    write_docx(out, "L", deck(), spans)
    assert headings(out) == [("Slide 1 – Intro", None), ("Slide 2", True), ("Slide 1 – Intro", True),
                             ("Slide 3 – Outro", True)]  # the first stays with the document title
    items = [(s, t) for s, t, _ in paragraphs(out)]
    assert [t for s, t in items if s == "Subtitle"] == ["00:00 – 00:10", "00:10 – 00:20", "00:20 – 00:30", "00:30 – 00:40"]
    texts = [t for s, t in items if s == "Normal" and t]
    assert texts == ["First visit.", "Middle.", "Back again.", "The end."]  # no [mm:ss] prefixes: the page has its time
    assert len(Document(out).inline_shapes) == 4  # one thumbnail per page


def test_pages_follow_the_lecture_not_the_slide_order_and_unsorted_input_is_sorted(tmp_path):
    out = tmp_path / "o.docx"
    spans = [SpanText(3, 30, "c", "c", end=40), SpanText(1, 0, "a", "a", end=10), SpanText(2, 10, "b", "b", end=30)]
    write_docx(out, "L", deck(), spans)
    assert [h for h, _ in headings(out)] == ["Slide 1 – Intro", "Slide 2", "Slide 3 – Outro"]
    spans = [SpanText(3, 0, "c", "c", end=10), SpanText(1, 10, "a", "a", end=20)]
    write_docx(out, "L", deck(), spans)
    assert [h for h, _ in headings(out)] == ["Slide 3 – Outro", "Slide 1 – Intro"]  # slide 3 was shown first


def test_slides_that_were_never_shown_get_no_page(tmp_path):
    out = tmp_path / "o.docx"
    write_docx(out, "L", deck(), [SpanText(2, 0, "Only this.", "r", end=5)])
    assert headings(out) == [("Slide 2", None)]


def test_a_segment_without_speech_says_so(tmp_path):
    out = tmp_path / "o.docx"
    write_docx(out, "L", deck(), [SpanText(1, 0, "   ", "", end=5), SpanText(2, 5, "Words.", "r", end=9)])
    items = [(s, t) for s, t, _ in paragraphs(out)]
    assert items.count(("Normal", "(no speech)")) == 1
    assert headings(out) == [("Slide 1 – Intro", None), ("Slide 2", True)]


def test_options_page_breaks_timestamps_thumbnails_and_raw(tmp_path):
    out = tmp_path / "o.docx"
    spans = [SpanText(1, 0, "Clean one.", "raw one", end=5), SpanText(2, 5, "Clean two.", "raw two", end=9)]
    write_docx(out, "L", deck(), spans, page_breaks=False, timestamps=False, thumbnails=False, keep_raw=True)
    assert headings(out) == [("Slide 1 – Intro", None), ("Slide 2", None)]
    items = [(s, t) for s, t, _ in paragraphs(out)]
    assert not [1 for s, _ in items if s == "Subtitle"]
    assert ("Normal", "Raw: raw one") in items and ("Normal", "Raw: raw two") in items
    assert len(Document(out).inline_shapes) == 0


def test_segments_for_slides_outside_the_deck_are_skipped_and_control_characters_removed(tmp_path):
    out = tmp_path / "o.docx"
    spans = [SpanText(99, 0, "Lost.", "r", end=5), SpanText(1, 5, "Te\x00xt\x0b.", "r\x1f", end=9)]
    write_docx(out, "L\x00", deck(), spans, keep_raw=True)
    items = [(s, t) for s, t, _ in paragraphs(out)]
    assert ("Title", "L") in items and ("Normal", "Text.") in items and ("Normal", "Raw: r") in items
    assert headings(out) == [("Slide 1 – Intro", None)]


def test_a_segment_without_a_known_end_has_no_time_line(tmp_path):
    out = tmp_path / "o.docx"
    write_docx(out, "L", deck(), [SpanText(1, 0, "Hi.", "r")])
    assert not [1 for s, _, _ in paragraphs(out) if s == "Subtitle"]
