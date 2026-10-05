import json
import re

import pytest
from docx import Document

import lecture2doc
from api import RateLimiter
from helpers import make_pdf, make_pptx, make_video
from slides import render_pdf
from transcribe import Word


def setup(tmp_path, monkeypatch):
    monkeypatch.setenv("SCADS_API_KEY", "dummy")
    pdf = make_pdf(tmp_path / "s.pdf", 3)
    pptx = make_pptx(tmp_path / "s.pptx", [("Alpha", "a", False), ("Beta", "b", False), ("Gamma", "c", False)])
    images, _ = render_pdf(pdf)
    video = make_video(images, [0] * 4 + [1] * 4 + [0] * 4, tmp_path / "lec.mp4")
    words = [Word("hello", 1.0, 1.5), Word("world", 5.0, 5.5), Word("again", 9.0, 9.5)]
    monkeypatch.setattr(lecture2doc, "transcribe_words", lambda *a, **k: words)
    return video, pdf, pptx


def test_end_to_end_no_cleanup(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    out = tmp_path / "out" / "lec.docx"
    assert lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(out), "--workdir", str(tmp_path / "w"),
                             "--no-cleanup"]) == 0
    texts = [(p.style.name, p.text) for p in Document(out).paragraphs]
    assert ("Heading 1", "Slide 1 – Alpha") in texts and ("Heading 1", "Slide 3 – Gamma") in texts
    assert ("Normal", "[00:00] hello") in texts
    # the revisit start time may differ by +-1 s from 8
    assert any(style == "Normal" and re.fullmatch(r"\[\d\d:\d\d\] again", t) for style, t in texts)
    assert ("Normal", "world") in texts
    assert ("Normal", "(no speech)") in texts  # slide 3 never shown
    seg = json.loads((tmp_path / "w" / "segments.json").read_text(encoding="utf-8"))
    assert [s["slide"] for s in seg["spans"]] == [1, 2, 1]
    assert "dummy" not in json.dumps(seg)


def test_from_segments_skips_matching(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    seg = tmp_path / "seg.json"
    seg.write_text(json.dumps({"spans": [{"slide": 3, "start": 0, "end": 12, "confidence": 1.0}]}))
    monkeypatch.setattr(lecture2doc, "extract_frames", lambda *a, **k: (_ for _ in ()).throw(AssertionError("matched")))
    out = tmp_path / "o.docx"
    lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(out), "--workdir", str(tmp_path / "w"),
                      "--no-cleanup", "--from-segments", str(seg)])
    texts = [p.text for p in Document(out).paragraphs]
    assert "hello world again" in texts


def test_from_segments_rejects_unknown_slide(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    seg = tmp_path / "seg.json"
    seg.write_text(json.dumps({"spans": [{"slide": 99, "start": 0, "end": 12, "confidence": 1.0}]}))
    with pytest.raises(SystemExit, match="99"):
        lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(tmp_path / "o.docx"),
                          "--workdir", str(tmp_path / "w"), "--no-cleanup", "--from-segments", str(seg)])


def run(video, pdf, pptx, tmp_path, *extra, out=None):
    out = out or tmp_path / "o.docx"
    rc = lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(out), "--workdir", str(tmp_path / "w"), *extra])
    assert rc == 0
    return out


def test_missing_video_fails_fast(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    with pytest.raises(SystemExit, match="video not found"):
        run(tmp_path / "nope.mp4", pdf, pptx, tmp_path, "--no-cleanup")


def test_missing_pdf_fails_fast(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    with pytest.raises(SystemExit, match="PDF not found"):
        run(video, tmp_path / "nope.pdf", pptx, tmp_path, "--no-cleanup")


def test_missing_segments_file_fails_fast(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    with pytest.raises(SystemExit, match="segments"):
        run(video, pdf, pptx, tmp_path, "--no-cleanup", "--from-segments", str(tmp_path / "nope.json"))


def test_from_segments_unsorted_matches_sorted(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    spans = [{"slide": 1, "start": 0, "end": 4, "confidence": 1.0},
             {"slide": 2, "start": 4, "end": 8, "confidence": 1.0},
             {"slide": 3, "start": 8, "end": 12, "confidence": 1.0}]
    results = []
    for name, order in (("sorted", spans), ("unsorted", spans[::-1])):
        seg = tmp_path / f"{name}.json"
        seg.write_text(json.dumps({"spans": order}))
        out = run(video, pdf, pptx, tmp_path, "--no-cleanup", "--from-segments", str(seg),
                  out=tmp_path / f"{name}.docx")
        results.append([(p.style.name, p.text) for p in Document(out).paragraphs])
    assert results[0] == results[1]
    assert [("Normal", w) for w in ("hello", "world", "again")] == [x for x in results[1] if x[1] in ("hello", "world", "again")]


@pytest.mark.parametrize("content, message", [
    ("{not json", "invalid JSON"),
    (json.dumps({"nospans": []}), "spans"),
    (json.dumps({"spans": "x"}), "spans"),
    (json.dumps({"spans": [{"slide": 1, "start": 0}]}), "end"),
    (json.dumps({"spans": [{"slide": 1, "start": "a", "end": 3}]}), "numeric"),
    (json.dumps({"spans": [{"slide": 1, "start": 5, "end": 5}]}), "end"),
    (json.dumps({"spans": [{"slide": 1, "start": 5, "end": 2}]}), "end"),
])
def test_from_segments_malformed(tmp_path, monkeypatch, content, message):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    seg = tmp_path / "seg.json"
    seg.write_text(content)
    with pytest.raises(SystemExit, match=message):
        run(video, pdf, pptx, tmp_path, "--no-cleanup", "--from-segments", str(seg))


def test_summary_per_slide(tmp_path, monkeypatch, capsys):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    run(video, pdf, pptx, tmp_path, "--no-cleanup")
    out = capsys.readouterr().out
    assert "slide  1:" in out and "slide  2:" in out
    assert "slide  3:" not in out
    assert "slides never shown: [3]" in out
    assert "unmatched" in out


def test_summary_for_from_segments(tmp_path, monkeypatch, capsys):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    seg = tmp_path / "seg.json"
    seg.write_text(json.dumps({"spans": [{"slide": 3, "start": 0, "end": 12, "confidence": 1.0}]}))
    run(video, pdf, pptx, tmp_path, "--no-cleanup", "--from-segments", str(seg))
    out = capsys.readouterr().out
    assert "slide  3:  12s matched" in out
    assert "slides never shown: [1, 2]" in out


def test_transcribe_runtime_error_exits(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)

    def boom(*a, **k):
        raise RuntimeError("no audio stream")

    monkeypatch.setattr(lecture2doc, "transcribe_words", boom)
    with pytest.raises(SystemExit, match="no audio stream"):
        run(video, pdf, pptx, tmp_path, "--no-cleanup")


def test_cleanup_wiring(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    monkeypatch.setenv("SCADS_LLM_MODEL", "llm-x")
    monkeypatch.setenv("SCADS_GLOSSARY_MODEL", "glos-y")
    monkeypatch.setenv("SCADS_REQUESTS_PER_MINUTE", "7")
    dummy = object()
    calls = {}
    monkeypatch.setattr(lecture2doc, "make_client", lambda cfg: dummy)

    def fake_glossary(client, model, slides, cache, limiter=None):
        calls["glossary"] = (client, model)
        calls["glossary_limiter"] = limiter
        return "GLOSS"

    def fake_cleanup(client, model, glossary, outline, slides, spans, raws, cache, limiter=None):
        calls["cleanup"] = (client, model, glossary)
        calls["cleanup_limiter"] = limiter
        return [f"CLEAN<{r}>" for r in raws]

    monkeypatch.setattr(lecture2doc, "build_glossary", fake_glossary)
    monkeypatch.setattr(lecture2doc, "cleanup_spans", fake_cleanup)
    out = run(video, pdf, pptx, tmp_path)
    assert calls["glossary"] == (dummy, "glos-y")
    assert calls["cleanup"] == (dummy, "llm-x", "GLOSS")
    assert isinstance(calls["cleanup_limiter"], RateLimiter) and calls["cleanup_limiter"].max_calls == 7
    assert calls["glossary_limiter"] is calls["cleanup_limiter"]  # one shared quota for both stages
    texts = [p.text for p in Document(out).paragraphs]
    assert "CLEAN<world>" in texts and "world" not in texts
    seg = json.loads((tmp_path / "w" / "segments.json").read_text(encoding="utf-8"))
    assert seg["spans"][1]["raw"] == "world"
    assert seg["spans"][1]["clean"] == "CLEAN<world>"


def test_keep_raw(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    out = run(video, pdf, pptx, tmp_path, "--no-cleanup", "--keep-raw")
    assert "Raw: world" in [p.text for p in Document(out).paragraphs]


def test_thumbnails_toggle(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    with_thumbs = run(video, pdf, pptx, tmp_path, "--no-cleanup", out=tmp_path / "a.docx")
    without = run(video, pdf, pptx, tmp_path, "--no-cleanup", "--no-thumbnails", out=tmp_path / "b.docx")
    assert len(Document(with_thumbs).inline_shapes) >= 1
    assert len(Document(without).inline_shapes) == 0


def test_probe_failure_falls_back_to_16_9_with_warning(tmp_path, monkeypatch, caplog):
    video, pdf, pptx = setup(tmp_path, monkeypatch)

    def bad_probe(v):
        raise RuntimeError("ffprobe could not read the video size")

    monkeypatch.setattr(lecture2doc, "probe_aspect", bad_probe)
    run(video, pdf, pptx, tmp_path, "--no-cleanup")
    assert "assuming 16:9" in caplog.text


def test_frame_extraction_error_exits(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)

    def boom(*a, **k):
        raise RuntimeError("no video frames extracted from lec.mp4")

    monkeypatch.setattr(lecture2doc, "extract_frames", boom)
    with pytest.raises(SystemExit, match="no video frames extracted"):
        run(video, pdf, pptx, tmp_path, "--no-cleanup")


def test_transcribe_api_error_exits_without_traceback_or_key(tmp_path, monkeypatch):
    import httpx2
    import openai

    video, pdf, pptx = setup(tmp_path, monkeypatch)

    def denied(*a, **k):
        raise openai.AuthenticationError(
            "invalid credentials",
            response=httpx2.Response(401, request=httpx2.Request("POST", "http://x")), body=None)

    monkeypatch.setattr(lecture2doc, "transcribe_words", denied)
    with pytest.raises(SystemExit) as exc_info:
        run(video, pdf, pptx, tmp_path, "--no-cleanup")
    msg = str(exc_info.value)
    assert "speech-to-text failed" in msg
    assert "invalid credentials" in msg
    assert "dummy" not in msg


def subtitle_lines(path):
    return [p.text for p in Document(path).paragraphs if p.style.name == "Subtitle"]


def non_subtitle(path):
    return [(p.style.name, p.text) for p in Document(path).paragraphs if p.style.name != "Subtitle"]


def test_normal_run_adds_time_subtitles(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    out = run(video, pdf, pptx, tmp_path, "--no-cleanup")
    lines = subtitle_lines(out)
    assert len(lines) == 2  # slide 1 (two visits) and slide 2; slide 3 never shown
    assert re.fullmatch(r"\d\d:\d\d – \d\d:\d\d · \d\d:\d\d – \d\d:\d\d", lines[0])
    assert re.fullmatch(r"\d\d:\d\d – \d\d:\d\d", lines[1])
    assert subtitle_lines(run(video, pdf, pptx, tmp_path, "--no-cleanup", "--no-timestamps",
                              out=tmp_path / "plain.docx")) == []


def forbid_heavy_work(monkeypatch):
    def boom(*a, **k):
        raise AssertionError("heavy work during --rebuild")
    for name in ("transcribe_words", "extract_frames", "match_frames", "probe_aspect", "load_config",
                 "check_ffmpeg", "make_client", "build_glossary", "cleanup_spans"):
        monkeypatch.setattr(lecture2doc, name, boom)


def test_rebuild_adds_subtitles_offline_without_video_or_key(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    old = run(video, pdf, pptx, tmp_path, "--no-cleanup", "--no-timestamps", out=tmp_path / "old.docx")
    assert subtitle_lines(old) == []
    video.unlink()  # the video is not needed any more
    forbid_heavy_work(monkeypatch)
    new = tmp_path / "new.docx"
    rc = lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(new), "--workdir", str(tmp_path / "w"),
                           "--rebuild"])
    assert rc == 0
    assert len(subtitle_lines(new)) == 2
    assert non_subtitle(new) == non_subtitle(old)  # identical text, only the time lines were added


def test_rebuild_uses_cleaned_text_and_keeps_segments_file_untouched(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    run(video, pdf, pptx, tmp_path, "--no-cleanup")
    segfile = tmp_path / "w" / "segments.json"
    data = json.loads(segfile.read_text(encoding="utf-8"))
    data["spans"][1]["clean"] = "World, corrected."
    segfile.write_text(json.dumps(data), encoding="utf-8")
    before = segfile.read_text(encoding="utf-8")
    forbid_heavy_work(monkeypatch)
    out = tmp_path / "r.docx"
    lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(out), "--workdir", str(tmp_path / "w"),
                      "--rebuild", "--keep-raw"])
    texts = [p.text for p in Document(out).paragraphs]
    assert "World, corrected." in texts and "Raw: world" in texts
    assert segfile.read_text(encoding="utf-8") == before


def test_rebuild_from_explicit_segments_file(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    seg = tmp_path / "elsewhere.json"
    seg.write_text(json.dumps({"spans": [{"slide": 2, "start": 5, "end": 65, "raw": "r", "clean": "Cleaned."}]}),
                   encoding="utf-8")
    forbid_heavy_work(monkeypatch)
    out = tmp_path / "e.docx"
    lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(out), "--workdir", str(tmp_path / "w2"),
                      "--rebuild", "--from-segments", str(seg)])
    assert subtitle_lines(out) == ["00:05 – 01:05"]


@pytest.mark.parametrize("content, message", [
    (None, "segments file"),                                                   # no file at all
    ({"spans": [{"slide": 1, "start": 0, "end": 5, "raw": "r"}]}, "'clean'"),  # old/hand-made file
    ({"spans": [{"slide": 1, "start": 0, "end": 5, "clean": 3, "raw": "r"}]}, "'clean'"),
])
def test_rebuild_requires_a_complete_segments_file(tmp_path, monkeypatch, content, message):
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    if content is not None:
        (tmp_path / "w").mkdir()
        (tmp_path / "w" / "segments.json").write_text(json.dumps(content), encoding="utf-8")
    forbid_heavy_work(monkeypatch)
    with pytest.raises(SystemExit, match=message):
        lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(tmp_path / "x.docx"),
                          "--workdir", str(tmp_path / "w"), "--rebuild"])


def test_page_breaks_by_default_and_no_page_breaks_option(tmp_path, monkeypatch):
    video, pdf, pptx = setup(tmp_path, monkeypatch)

    def flags(path):
        return [p.paragraph_format.page_break_before for p in Document(path).paragraphs if p.style.name == "Heading 1"]

    default = run(video, pdf, pptx, tmp_path, "--no-cleanup")
    assert flags(default) == [None, True, True]
    compact = run(video, pdf, pptx, tmp_path, "--no-cleanup", "--no-page-breaks", out=tmp_path / "compact.docx")
    assert flags(compact) == [None, None, None]
    forbid_heavy_work(monkeypatch)  # --rebuild honours the option too, offline
    rebuilt = tmp_path / "rebuilt.docx"
    lecture2doc.main([str(video), str(pdf), str(pptx), "-o", str(rebuilt), "--workdir", str(tmp_path / "w"),
                      "--rebuild", "--no-page-breaks"])
    assert flags(rebuilt) == [None, None, None]


def snap_setup(tmp_path, monkeypatch):
    """The lecturer clicks to the next slide (~t=4) while still finishing a sentence that ends at 4.6."""
    video, pdf, pptx = setup(tmp_path, monkeypatch)
    words = [Word("done.", 0.5, 0.9), Word("Then", 3.1, 3.3), Word("more", 3.4, 3.6), Word("words", 3.7, 3.9),
             Word("end.", 4.2, 4.6), Word("Next", 5.0, 5.4), Word("slide.", 5.6, 5.9)]
    monkeypatch.setattr(lecture2doc, "transcribe_words", lambda *a, **k: words)
    return video, pdf, pptx


def slide_paragraphs(path, number):
    """Text paragraphs under the heading of the given slide."""
    out, inside = [], False
    for p in Document(path).paragraphs:
        if p.style.name == "Heading 1":
            inside = p.text.startswith(f"Slide {number} ")
        elif inside and p.style.name == "Normal" and p.text:  # the thumbnail's paragraph has no text
            out.append(re.sub(r"^\[[\d:]+\] ", "", p.text))
    return out


def test_sentence_tail_stays_with_its_slide_by_default(tmp_path, monkeypatch):
    video, pdf, pptx = snap_setup(tmp_path, monkeypatch)
    out = run(video, pdf, pptx, tmp_path, "--no-cleanup")
    assert any(t.endswith("Then more words end.") for t in slide_paragraphs(out, 1))
    assert slide_paragraphs(out, 2) == ["Next slide."]
    seg = json.loads((tmp_path / "w" / "segments.json").read_text(encoding="utf-8"))
    first, second = seg["spans"][0], seg["spans"][1]
    assert first["end"] <= 5.0 and first["text_end"] > first["end"] - 1  # visual time stays on-screen time
    assert first["text_end"] == second["text_start"] and 4.6 < first["text_end"] < 5.0


def test_snap_seconds_zero_keeps_the_raw_cut(tmp_path, monkeypatch):
    video, pdf, pptx = snap_setup(tmp_path, monkeypatch)
    out = run(video, pdf, pptx, tmp_path, "--no-cleanup", "--snap-seconds", "0")
    assert not any("end." in t for t in slide_paragraphs(out, 1))
    assert slide_paragraphs(out, 2) == ["end. Next slide."]
